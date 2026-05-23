from __future__ import annotations

import io
from pathlib import Path

from .models import AnalysisReport


def report_to_docx_bytes(report: AnalysisReport) -> bytes:
    try:
        from docx import Document  # type: ignore
        from docx.shared import Pt  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "Word export requires python-docx. Install with: pip install python-docx"
        ) from exc

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(f"SDS case analysis report — {report.case_reference}", 0)
    doc.add_paragraph(report.summary)
    doc.add_paragraph(
        f"Readiness band: {report.readiness_band.value} ({report.readiness_score}/100)"
    )

    if report.model_prediction:
        doc.add_heading("ML conviction confidence (supportive only)", level=1)
        mp = report.model_prediction
        doc.add_paragraph(f"Label: {mp.confidence_label}")
        doc.add_paragraph(f"Probability: {mp.conviction_probability * 100:.1f}%")
        doc.add_paragraph(f"Model: {mp.model_name}")
        doc.add_paragraph("Human oversight required: Yes")

    doc.add_heading("NLP extracted features", level=1)
    for feat in report.extracted_features:
        doc.add_paragraph(f"{feat.label}: {feat.value:.2f} ({feat.source})", style="List Bullet")

    doc.add_heading("Weighted evidence scores", level=1)
    for ws in report.weighted_scores:
        doc.add_paragraph(
            f"{ws.category} (weight {ws.weight:.2f}): score {ws.score:.2f} -> {ws.weighted_contribution:+.2f}",
            style="List Bullet",
        )

    if report.missing_elements:
        doc.add_heading("Missing conviction elements", level=1)
        for gap in report.missing_elements:
            doc.add_paragraph(
                f"[{gap.severity.value.upper()}] {gap.element}: {gap.detail}",
                style="List Bullet",
            )

    if report.legal_cross_references:
        doc.add_heading("Trinidad & Tobago legal cross-references", level=1)
        for ref in report.legal_cross_references:
            doc.add_heading(f"{ref.citation} — {ref.section_ref}: {ref.title}", level=2)
            doc.add_paragraph(f"Source: {ref.source_title}")
            doc.add_paragraph(ref.summary)
            doc.add_paragraph(f"Relevance: {ref.relevance}")

    doc.add_heading("Documents analysed", level=1)
    for item in report.documents:
        types = ", ".join(item.detected_types)
        doc.add_paragraph(
            f"{item.filename} ({item.word_count} words) — {types}",
            style="List Bullet",
        )

    doc.add_heading("Strengths", level=1)
    for item in report.strengths:
        doc.add_paragraph(f"{item.title} ({item.category}): {item.detail}", style="List Bullet")

    doc.add_heading("Weaknesses and gaps", level=1)
    for item in report.weaknesses:
        doc.add_paragraph(
            f"[{item.severity.value.upper()}] {item.title} ({item.category}): {item.detail}",
            style="List Bullet",
        )

    doc.add_heading("Bulletproofing recommendations", level=1)
    for rec in report.recommendations:
        doc.add_paragraph(
            f"{rec.priority}. {rec.action} — {rec.rationale} (Ref: {rec.legal_anchor})"
        )

    doc.add_heading("Anticipated King's Counsel cross-examination", level=1)
    for idx, q in enumerate(report.cross_examination, 1):
        doc.add_heading(f"{idx}. {q.theme}", level=2)
        doc.add_paragraph(f"Target: {q.target_witness}")
        doc.add_paragraph(f"Question: {q.question}")
        doc.add_paragraph(f"Purpose: {q.purpose}")
        doc.add_paragraph(f"Follow-up: {q.follow_up}")

    doc.add_heading("Bias and fairness notice", level=1)
    doc.add_paragraph(report.bias_notice)
    doc.add_paragraph(report.disclaimer)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def write_report_docx(report: AnalysisReport, path: Path) -> None:
    path.write_bytes(report_to_docx_bytes(report))
